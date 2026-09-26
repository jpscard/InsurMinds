"""Gera o relatório do Projeto Final (.docx e .pdf, em docs/relatorio/) a partir de docs/relatorio/relatorio.md.

Mesmo visual do relatório do Desafio 5 (Segoe UI, títulos em azul-marinho, tabelas com cabeçalho
azul). O PDF é exportado pelo Microsoft Word (Windows); sem Word, só o .docx é gerado.

Execute:  python scripts/gerar_relatorio.py            (requer: pip install python-docx)
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "relatorio" / "relatorio.md"
OUT = ROOT / "docs" / "relatorio" / "InsurMinds – Projeto Final.docx"

FONT = "Segoe UI"
NAVY = RGBColor(0x0F, 0x2B, 0x48)
BLUE_DARK = RGBColor(0x1E, 0x3A, 0x8A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
TEXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x64, 0x74, 0x8B)
TEXT_WIDTH_IN = 6.5
HEADER_TXT = "Apólis — Projeto Final InsurMinds"
FOOTER_TXT = "InsurMinds — Projeto Final | Grupo JL (Leonardo Pereira & João Cardoso)"


# ─── Utilidades de formatação ───────────────────────────────────────────────
def shade(cell, fill: str) -> None:
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def borders(cell, top="E2E8F0", bottom="E2E8F0", left=None, right=None, sz="4") -> None:
    def side(name, color):
        return (f'<w:{name} w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>' if color
                else f'<w:{name} w:val="nil"/>')
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>{side("top", top)}{side("left", left)}{side("bottom", bottom)}{side("right", right)}</w:tcBorders>'))


def cell_margins(cell, v=70, h=100) -> None:
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{v}" w:type="dxa"/><w:bottom w:w="{v}" w:type="dxa"/>'
        f'<w:left w:w="{h}" w:type="dxa"/><w:right w:w="{h}" w:type="dxa"/></w:tcMar>'))


def style_run(run, size=10.0, color=TEXT, bold=False, italic=False, font=FONT) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def add_hyperlink(paragraph, text: str, url: str, size: float) -> None:
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:rFonts", {"w:ascii": FONT, "w:hAnsi": FONT}), ("w:color", {"w:val": "2563EB"}),
                       ("w:u", {"w:val": "single"}), ("w:sz", {"w:val": str(int(size * 2))})):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s)]+)")


def add_inline(paragraph, text: str, size=10.0, color=TEXT, italic=False, bold=False) -> None:
    """Markdown em linha: **negrito**, *itálico*, `código`, [texto](url) e URLs soltas."""
    for tok in _INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            add_inline(paragraph, tok[2:-2], size, color, italic, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            style_run(paragraph.add_run(tok[1:-1]), size - 0.5, RGBColor(0x0F, 0x17, 0x2A), bold, font="Consolas")
        elif tok.startswith("[") and "](" in tok:
            m = re.match(r"\[(.+?)\]\((.+?)\)", tok)
            add_hyperlink(paragraph, m.group(1), m.group(2), size)
        elif tok.startswith("http"):
            add_hyperlink(paragraph, tok, tok, size)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            style_run(paragraph.add_run(tok[1:-1]), size, color, bold, True)
        else:
            style_run(paragraph.add_run(tok), size, color, bold, italic)


def picture(doc_or_cell, path: Path, width_in: float, caption: str | None, center=True):
    p = doc_or_cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = bool(caption)
    p.add_run().add_picture(str(path), width=Inches(width_in))
    if caption:
        c = doc_or_cell.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        style_run(c.add_run(caption), 8.5, MUTED, italic=True)


# ─── Documento ──────────────────────────────────────────────────────────────
def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    sec.left_margin = sec.right_margin = Inches(1.0)
    for part, text, align in ((sec.header, HEADER_TXT, WD_ALIGN_PARAGRAPH.RIGHT),
                              (sec.footer, FOOTER_TXT, WD_ALIGN_PARAGRAPH.LEFT)):
        p = part.paragraphs[0]
        p.alignment = align
        style_run(p.add_run(text), 8.5, RGBColor(0x94, 0xA3, 0xB8))
    # número da página no rodapé, à direita
    p = sec.footer.paragraphs[0]
    p.add_run("\t\t")
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        style_run(r, 8.5, RGBColor(0x94, 0xA3, 0xB8))
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    for name, size, color, before, after in (("Heading 1", 14, NAVY, 18, 6), ("Heading 2", 12, BLUE_DARK, 12, 4),
                                             ("Heading 3", 11, BLUE, 10, 3)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = FONT, Pt(size), True, color
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:asciiTheme"), "")
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
        st.paragraph_format.keep_with_next = True


def title_block(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    style_run(p.add_run("RELATÓRIO TÉCNICO — INSURMINDS PROJETO FINAL"), 12, BLUE, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(14)
    style_run(p2.add_run(title), 16, NAVY, bold=True)


def _col_widths(rows: list[list[str]]) -> list[float]:
    """Largura proporcional ao texto, sem quebrar palavras: cada coluna cabe a maior palavra."""
    plain = [[re.sub(r"[*`]", "", c) for c in r] for r in rows]
    ncols = len(rows[0])
    minimo = [max(len(w) for r in plain for w in (r[i].split() or [""])) * 0.075 + 0.22 for i in range(ncols)]
    if sum(minimo) >= TEXT_WIDTH_IN:  # nem os mínimos cabem: reduz todos na mesma proporção
        return [m * TEXT_WIDTH_IN / sum(minimo) for m in minimo]
    peso = [sum(len(r[i]) for r in plain) / len(plain) + 4 for i in range(ncols)]
    widths = [max(minimo[i], TEXT_WIDTH_IN * peso[i] / sum(peso)) for i in range(ncols)]
    excesso = sum(widths) - TEXT_WIDTH_IN
    if excesso > 0:  # tira só das colunas acima do mínimo, sem passar dele
        folga = [w - m for w, m in zip(widths, minimo)]
        widths = [w - excesso * f / sum(folga) for w, f in zip(widths, folga)]
    return widths


def md_table(doc: Document, rows: list[list[str]], base: Path) -> None:
    is_image = any("![" in c for r in rows for c in r)
    if is_image:  # a legenda vem do texto alternativo de cada imagem; a linha de rótulos sobraria
        rows = [r for r in rows if any("![" in c for c in r)]
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [TEXT_WIDTH_IN / ncols] * ncols if is_image else _col_widths(rows)
    for ri, row in enumerate(rows):
        header = ri == 0 and not is_image
        for ci, text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            if is_image:
                borders(cell, None, None)
            elif header:
                shade(cell, "1E3A8A")
                borders(cell, "1E3A8A", "1E3A8A")
            else:
                shade(cell, "F8FAFC" if ri % 2 else "FFFFFF")
                borders(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            m = re.search(r"!\[(.*?)\]\((.*?)\)", text)
            if m:
                cell._tc.remove(p._p)
                picture(cell, base / m.group(2), widths[ci] - 0.2, m.group(1))
            elif header:
                style_run(p.add_run(text), 9, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
            else:
                add_inline(p, text, 8.5)
    for row in table.rows:  # não quebrar uma linha entre páginas
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(TEXT_WIDTH_IN)
    shade(cell, "F1F5FE")
    borders(cell, "DBEAFE", "DBEAFE", "2563EB", "DBEAFE", sz="8")
    cell_margins(cell, 100, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    add_inline(p, text, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(src: Path, out: Path) -> Path:
    base = src.parent
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup(doc)
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("# "):
            title_block(doc, s[2:].strip())
        elif s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=2)
        elif s.startswith(">"):
            buf = [s.lstrip("> ")]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith(">"):
                i += 1
                buf.append(lines[i].strip().lstrip("> "))
            callout(doc, " ".join(buf))
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                r = lines[i].strip()
                if not re.match(r"^\|(\s*:?-+:?\s*\|)+$", r):
                    rows.append([c.strip() for c in r.strip("|").split("|")])
                i += 1
            md_table(doc, rows, base)
            continue
        elif (m := re.match(r"^!\[(.*?)\]\((.*?)\)$", s)):
            picture(doc, base / m.group(2), TEXT_WIDTH_IN, f"Figura: {m.group(1)}")
        elif (m := re.match(r"^[-*]\s+(.*)$", s)):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(1), 9.5)
        elif (m := re.match(r"^(\d+)\.\s+(.*)$", s)):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.space_after = Pt(3)
            style_run(p.add_run(f"{m.group(1)}. "), 9.5, BLUE_DARK, bold=True)
            add_inline(p, m.group(2), 9.5)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, s, 10)
        i += 1
    doc.core_properties.title = "Apólis — Relatório Técnico do Projeto Final"
    doc.core_properties.author = "Grupo JL (Leonardo Pereira & João Cardoso)"
    doc.save(out)
    return out


def to_pdf(docx: Path) -> Path | None:
    """Exporta para PDF com o Microsoft Word (mesmo método do Desafio 5)."""
    pdf = docx.with_suffix(".pdf")
    ps = f"""
$w = New-Object -ComObject Word.Application; $w.Visible = $false
try {{ $d = $w.Documents.Open('{docx}'); $d.SaveAs([ref]'{pdf}', [ref]17); $d.Close() }}
finally {{ $w.Quit() }}"""
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, capture_output=True, timeout=180)
    except (OSError, subprocess.SubprocessError) as exc:
        print(f"PDF não gerado (Word indisponível?): {exc}")
        return None
    return pdf if pdf.exists() else None


if __name__ == "__main__":
    out = build(SRC, OUT)
    print(f"DOCX: {out}")
    if "--sem-pdf" not in sys.argv:
        pdf = to_pdf(out)
        if pdf:
            print(f"PDF:  {pdf}")
